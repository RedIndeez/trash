# КС-2 / КС-3 по образцу акта №1 (лит. 4) на весь объём ВДЦ, объём = 70 % площади по договору.
# Плюс журнал КС-6а и лист согласования ИД по образцам.
# Запуск: python tools/ks.py <ВДЦ.xlsx> <образец КС.xlsx> <образец КС-6а.xlsx> <лист согласования.xlsx>
import sys, re, copy, warnings
import openpyxl
from openpyxl.formula.translate import Translator
warnings.filterwarnings('ignore')

VDC_F, KS_F, J6_F, LS_F = sys.argv[1:5]
K = 0.7                       # доля договорной площади к закрытию
ACT_NO, ACT_DATE = 1, '25.09.2026г.'

# ---------------------------------------------------------------- ВДЦ
vws = openpyxl.load_workbook(VDC_F).worksheets[0]
SECTIONS = []                 # [(название, [группа, [(текст работ, текст изделия, площадь, работа, материал)]])]
sec = grp = None
for r in range(16, 146):
    a, b = vws.cell(r, 1).value, vws.cell(r, 2).value
    if b is None and isinstance(a, str) and a.strip():
        t = a.strip()
        if t.startswith('Блок-секция'):
            sec = (t, []); SECTIONS.append(sec); grp = None
        else:
            grp = (t, []); sec[1].append(grp)
        continue
    if isinstance(b, str) and b.startswith('Изготовление'):
        if grp is None:
            grp = ('Витражные блоки', []); sec[1].append(grp)
        e = vws.cell(r, 5).value
        area = round(eval(e.lstrip('=')) if isinstance(e, str) else e, 2)
        item = re.sub(r'\s*-\s*\d+\s*шт\.?\s*$', '', vws.cell(r + 1, 2).value.strip())
        grp[1].append((b, item, area, vws.cell(r, 7).value, vws.cell(r + 1, 8).value))

# ---------------------------------------------------------------- образец КС-2
wb = openpyxl.load_workbook(KS_F)
ws = wb['КС-2 ']
MAXC = ws.max_column


def snap(r):
    cells = {}
    for c in range(1, MAXC + 1):
        x = ws.cell(r, c)
        cells[c] = (x.value, copy.copy(x.font), copy.copy(x.border), copy.copy(x.fill), x.number_format,
                    copy.copy(x.alignment), copy.copy(x.protection))
    merges = [(m.min_col, m.max_col, m.max_row - m.min_row) for m in ws.merged_cells.ranges if m.min_row == r]
    return dict(cells=cells, h=ws.row_dimensions[r].height, merges=merges, src=r)


T_SEC, T_GRP, T_POS1, T_POS2, T_TOT = snap(36), snap(37), snap(38), snap(39), snap(51)
TAIL = [snap(r) for r in range(120, ws.max_row + 1)]
TAIL0 = 120

for m in [m for m in ws.merged_cells.ranges if m.min_row >= 36]:
    ws.unmerge_cells(str(m))
for r in range(36, ws.max_row + 1):
    ws.row_dimensions[r].height = None
    for c in range(1, MAXC + 1):
        x = ws.cell(r, c); x.value = None; x.style = 'Normal'


def paste(s, r, values=None, translate=False):
    for c, (v, f, b, fi, nf, al, pr) in s['cells'].items():
        x = ws.cell(r, c)
        if translate and isinstance(v, str) and v.startswith('='):
            v = Translator(v, origin=f'A{s["src"]}').translate_formula(f'A{r}')
        x.value = None if values is not None and not translate else v
        x.font, x.border, x.fill, x.number_format, x.alignment, x.protection = f, b, fi, nf, al, pr
    for c, v in (values or {}).items():
        ws[f'{c}{r}'] = v
    ws.row_dimensions[r].height = s['h']
    for c0, c1, dr in s['merges']:
        ws.merge_cells(start_row=r, start_column=c0, end_row=r + dr, end_column=c1)


r = 36
sec_tot = []
for sname, groups in SECTIONS:
    paste(T_SEC, r, {'C': sname}); r += 1
    first = r; n = 1
    for gname, items in groups:
        paste(T_GRP, r, {'B': gname}); r += 1
        for work, item, area, pw, pm in items:
            # вторая строка первой: объединения первой строки (N:U, V:Z, AC:AE) захватывают и её
            paste(T_POS2, r + 1, {'C': item, 'L': 'м2', 'AB': pm, 'AG': f'=AB{r+1}*N{r}'})
            paste(T_POS1, r, {'A': n, 'B': work, 'N': f'=ROUND({area}*{K},2)', 'V': f'=AA{r}+AB{r+1}', 'AA': pw,
                              'AC': f'=AF{r}+AG{r+1}', 'AF': f'=N{r}*AA{r}'})
            if len(work) > 300: ws.row_dimensions[r].height = 80   # в образце 63 — длинное описание витража обрезается
            r += 2; n += 1
    paste(T_TOT, r, {'L': 'м2', 'R': f'=SUM(N{first}:U{r-1})', 'AC': f'=SUM(AC{first}:AE{r-1})'})
    sec_tot.append(r); r += 1

off = r - TAIL0
for s in TAIL:
    paste(s, s['src'] + off, translate=True)
tot = TAIL0 + off
ws[f'AB{tot}'] = '=' + '+'.join(f'R{x}' for x in sec_tot)
ws[f'AG{tot}'] = '=' + '+'.join(f'AC{x}' for x in sec_tot)
ws['O29'] = f'=AG{tot+3}'
ws['N24'] = ACT_NO; ws['S24'] = ACT_DATE
ws.print_area = f'A1:AG{ws.max_row - 2}'

# ---------------------------------------------------------------- КС-3
k3 = wb['КС-3']
k3['I34'] = f"='КС-2 '!AG{tot}"
k3['F27'] = '=I32'                           # акт №1: с начала работ = за период (без НДС, как гр. 6)
k3['E20'] = str(ACT_NO); k3['F20'] = ACT_DATE

# спецификация по КП из образца к этому акту не относится
if 'Спецификация' in wb.sheetnames:
    del wb['Спецификация']
wb._external_links = []       # ссылка на внешний файл жила только в удалённом листе — иначе Excel «восстанавливает» книгу
wb.calculation.fullCalcOnLoad = True
wb.save('КС-2_КС-3_№1_витражи_лит4_70%_ВДЦ.xlsx')

# ---------------------------------------------------------------- КС-6а
jb = openpyxl.load_workbook(J6_F); j = jb.active
for rr in range(40, j.max_row + 1):
    if j.cell(rr, 6).value == 2966:          # строка «Изготовление и монтаж»
        j.cell(rr, 9).value = f'=ROUND(ROUND(D{rr},2)*{K},2)'
        j.cell(rr, 10).value = f'=I{rr}*E{rr}'
        j.cell(rr, 11).value = f'=J{rr}'
        j.cell(rr, 12).value = f'=I{rr}'
        j.cell(rr, 13).value = f'=J{rr}'
jb.calculation.fullCalcOnLoad = True
jb.save('КС-6а_журнал_лит4.xlsx')

# ---------------------------------------------------------------- лист согласования ИД
lb = openpyxl.load_workbook(LS_F); l = lb.active
l['C6'] = f'КС-2 №{ACT_NO}  от {ACT_DATE}'
l['B11'] = 'Изготовление и монтаж алюминиевых дверей и витражей (секции: А, Б, В, Г, Д, Ж)'
lb.save('Лист_согласования_ИД_КС-2_№1.xlsx')

n_pos = sum(len(i) for _, g in SECTIONS for _, i in g)
print('позиций', n_pos, 'итог в строке', tot, 'секции', [s for s, _ in SECTIONS])
