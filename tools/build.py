import json, collections, openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter as L
inst=json.load(open('inst.json'))
H48=4820
S={'Вн-1':(4750,3870),'Вн-2':(4750,3870),'Вн-3':(6100,3870),'Вн-4':(5400,3470),'Вн-5':(7800,3470),
'Вн-6':(4200,H48),'Вн-7':(3020,H48),'Вн-8':(2600,H48),'Вн-9':(3870,H48),'Вн-10':(4160,H48),'Вн-11':(2450,H48),
'Вн-12':(3370,H48),'Вн-13':(3200,H48),'Вн-14':(2485,H48),'Вн-15':(1400,H48),'Вн-16':(1200,H48),'Вн-17':(4200,H48),
'Вн-18':(2600,H48),'Вн-19':(3000,H48),'Вн-20':(3900,H48),'Вн-21':(2920,H48),'Вн-22':(2150,2720),'Вн-23':(3540,3020),'Вн-24':(2610,3020)}
LITS=['А','Б','В','Г','Д']
key=lambda m:int(m.split('-')[1])
drw=collections.Counter((i['mark'],i['lit']) for i in inst)
# количество по КП (колонка «по КП или Изготовлено» таблицы заказчика)
KP={('Вн-6','А'):4,('Вн-8','А'):1,('Вн-14','А'):1,('Вн-15','А'):1,('Вн-22','А'):5,('Вн-23','А'):0,('Вн-24','А'):0,
('Вн-6','Б'):4,('Вн-15','Б'):1,('Вн-17','Б'):1,('Вн-18','Б'):1,('Вн-22','Б'):5,('Вн-23','Б'):0,('Вн-24','Б'):0,
('Вн-6','В'):1,('Вн-13','В'):1,('Вн-16','В'):3,('Вн-17','В'):1,('Вн-19','В'):2,('Вн-22','В'):1,
('Вн-6','Г'):1,('Вн-10','Г'):1,('Вн-11','Г'):1,('Вн-12','Г'):1,('Вн-17','Г'):1,('Вн-20','Г'):1,('Вн-21','Г'):1,('Вн-22','Г'):0,
('Вн-6','Д'):2,('Вн-7','Д'):1,('Вн-9','Д'):1,('Вн-17','Д'):1,('Вн-10','Д'):1,
('Вн-1','Ж'):0,('Вн-2','Ж'):0,('Вн-3','Ж'):0,('Вн-4','Ж'):0,('Вн-5','Ж'):0}
# ---- позиции КП (файл КП, 06.04.2026 / 25.09.2026): (КП, №, марка, ширина, высота, шт, площадь в КП всего, куда отнесено)
KPP=[
('Блок А 06.04',1,'Вн-6',4150,4800,3,48.63,'А'),('Блок А 06.04',2,'Вн-8',2570,4800,1,12.08,'А'),
('Блок А 06.04',3,'Вн-15',1170,4800,1,2.64,'А'),('Блок А 06.04',4,'Вн-14',2320,4800,1,7.55,'А'),
('Блок Б 06.04',1,'Вн-6',4040,4800,1,15.69,'Б'),('Блок Б 06.04',2,'Вн-17',4060,4800,1,19.08,'Б'),
('Блок Б 06.04',3,'Вн-6',4060,4800,1,15.79,'Б'),('Блок Б 06.04',4,'Вн-15',1250,4800,1,2.83,'Б'),
('Блок Б 06.04',5,'Вн-18',2480,4800,1,8.30,'Б'),('Блок Б 06.04',6,'Вн-6',4080,4800,1,15.88,'Б'),
('КП без шапки 06.04 (лит. Д)',1,'Вн-10',4200,4800,1,16.45,'не вошло'),('КП без шапки 06.04 (лит. Д)',2,'Вн-9',4100,4800,1,15.98,'Д'),
('КП без шапки 06.04 (лит. Д)',3,'Вн-17',3780,4800,1,17.77,'Д'),('КП без шапки 06.04 (лит. Д)',4,'Вн-6',4140,4800,1,16.16,'Д'),
('КП без шапки 06.04 (лит. Д)',5,'Вн-7',4160,4800,1,19.55,'Д'),('КП без шапки 06.04 (лит. Д)',6,'Вн-6',4190,4800,1,16.40,'Д'),
('Блок В 06.04 (1)',1,'Вн-16',1370,5200,1,3.10,'В'),('Блок В 06.04 (1)',2,'Вн-16',1060,5180,1,2.40,'В'),
('Блок В 06.04 (2)',1,'Вн-6',4090,4800,1,15.93,'В'),('Блок В 06.04 (2)',2,'Вн-17',4100,4800,1,19.27,'В'),
('Блок В 06.04 (2)',3,'Вн-16',1240,4800,1,2.80,'В'),('Блок В 06.04 (2)',4,'Вн-19',2570,4800,2,17.44,'В'),
('Вн-22 06.04',1,'Вн-22',2070,2760,5,28.71*5/11,'А'),('Вн-22 06.04',1,'Вн-22',2070,2760,5,28.71*5/11,'Б'),('Вн-22 06.04',1,'Вн-22',2070,2760,1,28.71/11,'В'),
('Режиссёр 1-18 25.09',1,'Вн-19',2860,4800,1,10.09,'не вошло'),('Режиссёр 1-18 25.09',2,'Вн-19',2920,4800,1,10.37,'не вошло'),
('Режиссёр 1-18 25.09',3,'Вн-13',3070,4800,1,14.43,'В'),('Режиссёр 1-18 25.09',4,'Вн-16',1320,4800,1,2.98,'не вошло'),
('Режиссёр 1-18 25.09',5,'Вн-16',1020,4800,1,2.31,'не вошло'),('Режиссёр 1-18 25.09',6,'Вн-6',4130,4800,1,16.12,'Г←А'),
('Режиссёр 1-18 25.09',7,'Вн-17',4130,4800,1,19.41,'Г'),('Режиссёр 1-18 25.09',8,'Вн-6',4450,4800,1,17.62,'Г←Б'),
('Режиссёр 1-18 25.09',9,'Вн-17',4020,4800,1,18.89,'не вошло'),('Режиссёр 1-18 25.09',10,'Вн-21',2920,4800,1,13.72,'Г'),
('Режиссёр 1-18 25.09',11,'Вн-6',4260,4800,1,16.73,'Г'),('Режиссёр 1-18 25.09',12,'Вн-20',3840,4800,1,18.05,'Г'),
('Режиссёр 1-18 25.09',13,'Вн-12',3250,4800,1,11.92,'Г'),('Режиссёр 1-18 25.09',14,'Вн-11',2310,4800,1,10.86,'Г'),
('Режиссёр 1-18 25.09',15,'Вн-10',4160,4800,1,16.26,'Г'),('Режиссёр 1-18 25.09',16,'Вн-9',4060,4800,1,15.79,'не вошло'),
('Режиссёр 1-18 25.09',17,'Вн-17',3750,4800,1,17.62,'не вошло'),('Режиссёр 1-18 25.09',18,'Вн-6',4130,4800,1,16.12,'не вошло'),
]
# двери к витражам, КП от 13.04.2026 (в КП витражей площадь дана без дверей)
KPD=[('№9 А','А','Вн-6',1336,2433,3,9.75),('№9 А','А','Вн-14',1336,2433,1,3.25),('№9 А','А','Вн-15',1106,2433,1,2.69),
('№11 Б','Б','Вн-6',1336,2433,3,9.75),('№11 Б','Б','Вн-18',1336,2433,1,3.25),('№11 Б','Б','Вн-15',1186,2433,1,2.89),
('№14 (Д)','Д','Вн-10',1336,2433,1,3.25),('№14 (Д)','Д','Вн-6',1336,2433,2,6.50),('№14 (Д)','Д','Вн-9',1336,2433,1,3.25),
('№12 В','В','Вн-6',1336,2433,1,3.25),('№12 В','В','Вн-19',1336,2433,2,6.50),('№12 В','В','Вн-16',1176,2433,1,2.86),
('№15 Вн-22','А, Б, В','Вн-22',1336,2193,11,32.23)]
gab=lambda w,h,q: round(w*h*q/1e6,2)
# Вн-6: на чертеже лит. Г — 3 шт., в КП/ВДЦ лит. Г — 1. Две шт. (КП «Режиссёр 1-18» №6 и №8) установлены в лит. Г,
# но записываются в объёмы лит. А и Б (там в КП и ВДЦ по 4 шт., на чертеже по 3)
OVR={('Вн-6','А'):4,('Вн-6','Б'):4,('Вн-6','Г'):1}
# площади по КП (колонка «по КП или Изготовлено», площадь), с учётом переноса Вн-6
MOVED={'А':gab(4130,4800,1),'Б':gab(4450,4800,1)}   # КП «Режиссёр 1-18» №6 и №8
PREV={('Вн-6','А'):59.68,('Вн-8','А'):12.34,('Вн-14','А'):11.14,('Вн-15','А'):5.62,('Вн-22','А'):28.57,
('Вн-6','Б'):59.86,('Вн-15','Б'):6.00,('Вн-17','Б'):19.49,('Вн-18','Б'):11.90,('Вн-22','Б'):28.57,
('Вн-6','В'):19.63,('Вн-13','В'):14.74,('Вн-16','В'):18.57,('Вн-17','В'):19.68,('Вн-19','В'):24.67,('Вн-22','В'):5.71,
('Вн-6','Г'):60.31,('Вн-10','Г'):19.97,('Вн-11','Г'):11.09,('Вн-12','Г'):15.60,('Вн-17','Г'):19.82,('Вн-20','Г'):18.43,('Вн-21','Г'):14.02,
('Вн-6','Д'):39.98,('Вн-7','Д'):19.97,('Вн-9','Д'):19.68,('Вн-17','Д'):18.14,('Вн-10','Д'):20.16}
KPA={}
for doc,p,m,w,h,q,a,l in KPP:
    if l in LITS or l in ('Г←А','Г←Б'): KPA[(m,l if l in LITS else 'Г')]=KPA.get((m,l),0)+w*h*q/1e6
KPA={k:round(v+1e-9,2) for k,v in KPA.items()}
KPA[('Вн-10','Д')]=gab(4200,4800,1)
KPA[('Вн-6','Г')]=gab(4260,4800,1)
def area(m,l):
    if (m,l)==('Вн-6','А'): return round(KPA[(m,l)]+MOVED['А'],2)
    if (m,l)==('Вн-6','Б'): return round(KPA[(m,l)]+MOVED['Б'],2)
    return KPA[(m,l)]
cnt=collections.Counter({k:min(v,KP[k]) for k,v in drw.items() if KP[k]>0})
cnt.update({k:v-cnt[k] for k,v in OVR.items()})
inst=[i for i in inst if cnt[(i['mark'],i['lit'])] or (i['mark'],i['lit'])==('Вн-6','Г')]
marks=sorted({i['mark'] for i in inst},key=key)
lvl={}; drs=collections.defaultdict(set)
for i in inst:
    lvl.setdefault(i['mark'],set()).add(i['file'].replace('План на отм. ',''))
    if i['door']: drs[i['mark']].add(i['door'])
NOTE={'Вн-22':'КП: 2070×2760. На чертеже 18 шт., в КП 11 (А5, Б5, В1)','Вн-6':'На чертеже А3, Б3, Г3. Две шт. из лит. Г (КП 1-18 №6, №8) записаны в объёмы лит. А и Б — по ВДЦ А4, Б4, Г1','Вн-17':'Лит. Г: на чертеже 2 шт., в КП 1',
      'Вн-16':'1 шт. у лестницы на границе литеров В/Г отнесён к литеру В'}
thin=Side(style='thin'); B=Border(left=thin,right=thin,top=thin,bottom=thin)
hf=Font(bold=True); fill=PatternFill('solid',fgColor='DDEBF7'); tf=PatternFill('solid',fgColor='F2F2F2')
C=Alignment(horizontal='center',vertical='center',wrap_text=True)
W=Alignment(vertical='center',wrap_text=True)
def box(ws,r1,r2,c1,c2):
    for r in range(r1,r2+1):
        for c in range(c1,c2+1): ws.cell(r,c).border=B

wb=openpyxl.Workbook()
# ---- Sheet 1
ws=wb.active; ws.title='Спецификация'
ws['A1']='Спецификация витражей (по планам на отм. -6.000 и 0.000, марки выделены красным) — только позиции, входящие в КП'; ws['A1'].font=Font(bold=True,size=12)
hdr=['№','Марка','Длина, мм','Высота, мм','Площадь 1 шт по КП, м²']+[f'Литер {l}' for l in LITS]+['Всего, шт','Площадь по КП, м²','Отметка','Дверь в составе (по чертежу)','Примечание']
ws.append([]); ws.append(hdr)
for c in range(1,len(hdr)+1): ws.cell(3,c).font=hf; ws.cell(3,c).alignment=C; ws.cell(3,c).fill=fill
r=4
for n,m in enumerate(marks,1):
    w,h=S[m]
    row=[n,m,w,h,f'=ROUND(L{r}/K{r},2)']+[cnt[(m,l)] or None for l in LITS]
    row+=[f'=SUM(F{r}:J{r})',round(sum(area(m,l) for l in LITS if cnt[(m,l)]),2),', '.join(sorted(lvl[m])),', '.join(sorted(drs[m])),NOTE.get(m,'')]
    ws.append(row); r+=1
last=r-1
ws.cell(r,2,'ИТОГО').font=hf
for c in range(6,13):
    col=L(c); ws.cell(r,c,f'=SUM({col}4:{col}{last})').font=hf
for c in range(1,len(hdr)+1): ws.cell(r,c).fill=tf
box(ws,3,r,1,len(hdr))
for rr in range(4,r+1):
    for c in range(1,len(hdr)+1):
        ws.cell(rr,c).alignment=W if c in(14,15) else C
    ws.cell(rr,5).number_format='0.00'; ws.cell(rr,12).number_format='#,##0.00'
ws.cell(r+2,1,'Размеры — из таблицы заказчика (длина × высота). Количество — по чертежам, но не больше, чем в КП. Площади — по КП (площадь 1 шт — средняя).')
for i,wd in enumerate([4,8,9,9,10,7,7,7,7,7,8,11,10,13,34],1): ws.column_dimensions[L(i)].width=wd
ws.row_dimensions[3].height=62
ws.freeze_panes='C4'
ws.page_setup.orientation='landscape'; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0; ws.sheet_properties.pageSetUpPr.fitToPage=True

# ---- Sheet 2
ws2=wb.create_sheet('По литерам')
h2=['№','Марка','Длина, мм','Высота, мм','Площадь 1 шт по КП, м²','Кол-во, шт','Площадь по КП, м²','Отметка']
ws2.append(['Спецификация витражей по литерам']); ws2['A1'].font=Font(bold=True,size=12)
ws2.append([]); ws2.append(h2)
for c in range(1,9): ws2.cell(3,c).font=hf; ws2.cell(3,c).alignment=C; ws2.cell(3,c).fill=fill
r=4; n=1; subs=[]
for l in LITS:
    ms=[m for m in marks if cnt[(m,l)]]
    ws2.cell(r,2,f'Литер {l}').font=hf; ws2.merge_cells(start_row=r,start_column=2,end_row=r,end_column=8); r+=1
    s=r
    for m in ms:
        w,h=S[m]; lv=', '.join(sorted({i['file'].replace('План на отм. ','') for i in inst if i['mark']==m and i['lit']==l}))
        if (m,l) in (('Вн-6','А'),('Вн-6','Б')):
            parts=[(m,3,KPA[(m,l)]),('Вн-6 (устан. в лит. Г, КП 1-18 №'+('6' if l=='А' else '8')+')',1,MOVED[l])]
        else: parts=[(m,cnt[(m,l)],area(m,l))]
        for nm,q,a in parts:
            ws2.append([n,nm,w,h,f'=ROUND(G{r}/F{r},2)',q,a,lv]); n+=1; r+=1
    ws2.cell(r,2,f'Итого литер {l}').font=hf
    ws2.cell(r,6,f'=SUM(F{s}:F{r-1})').font=hf; ws2.cell(r,7,f'=SUM(G{s}:G{r-1})').font=hf
    for c in range(1,9): ws2.cell(r,c).fill=tf
    subs.append(r); r+=1
ws2.cell(r,2,'ВСЕГО').font=hf
ws2.cell(r,6,'='+'+'.join(f'F{x}' for x in subs)).font=hf; ws2.cell(r,7,'='+'+'.join(f'G{x}' for x in subs)).font=hf
box(ws2,3,r,1,8)
for rr in range(4,r+1):
    for c in range(1,9):
        if ws2.cell(rr,c).coordinate not in ws2.merged_cells: ws2.cell(rr,c).alignment=C
    ws2.cell(rr,5).number_format='0.00'; ws2.cell(rr,7).number_format='#,##0.00'
for i,wd in enumerate([5,21,10,10,11,10,13,11],1): ws2.column_dimensions[L(i)].width=wd
ws2.row_dimensions[3].height=48
# ---- Sheet 3
ws3=wb.create_sheet('Привязка к чертежам')
h3=['№','Чертёж','Литер','Марка','Дверь','X, мм','Y, мм','В КП']
ws3.append(['Витражи на чертежах, входящие в КП (координаты — точка выноски, мм)'])
ws3.append(['«1 из 3» — на чертеже больше, чем в КП; какие именно — по КП не определить']); ws3['A1'].font=hf
ws3.append(h3)
for c in range(1,9): ws3.cell(3,c).font=hf; ws3.cell(3,c).alignment=C; ws3.cell(3,c).fill=fill
for n,i in enumerate(sorted(inst,key=lambda i:(i['file'],LITS.index(i['lit']),key(i['mark']))),1):
    k=(i['mark'],i['lit']); ws3.append([n,i['file'],i['lit'],i['mark'],i['door'],i['x'],i['y'],('1 по лит. Г, 2 записаны в лит. А и Б' if k==('Вн-6','Г') else 'да' if drw[k]<=KP[k] else f'{KP[k]} из {drw[k]}')])
box(ws3,3,ws3.max_row,1,8)
for rr in range(4,ws3.max_row+1):
    for c in range(1,9): ws3.cell(rr,c).alignment=C
    ws3.cell(rr,6).number_format='#,##0'; ws3.cell(rr,7).number_format='#,##0'
for i,wd in enumerate([5,20,7,9,8,11,11,9],1): ws3.column_dimensions[L(i)].width=wd
# ---- Sheet 4: есть в КП, но не вошло в спецификацию
ws4=wb.create_sheet('В КП, не вошло')
ws4.append(['Позиции КП, не вошедшие в спецификацию (нет на чертежах или на чертежах меньше, чем в КП)']); ws4['A1'].font=Font(bold=True,size=12)
ws4.append([])
h4=['№','Литер','Марка','Длина, мм','Высота, мм','Площадь 1 шт по КП, м²','В КП, шт','Вошло в спецификацию, шт','Не вошло, шт','Площадь не вошедшего, м²','Причина']
ws4.append(h4)
for c in range(1,12): ws4.cell(3,c).font=hf; ws4.cell(3,c).alignment=C; ws4.cell(3,c).fill=fill
rows=[]
for (m,l),q in KP.items():
    got=cnt[(m,l)]
    if q>got:
        w,h=S[m]
        why='нет на чертежах' if drw[(m,l)]==0 else f'на чертежах {drw[(m,l)]} шт.'
        rows.append((l,m,w,h,q,got,why,round(KPA[(m,l)]/q,2) if (m,l) in KPA else None))
# остаток КП без литера (в таблице заказчика — строки «без литера»)
for m,w in [('Вн-19',2860),('Вн-19',2920),('Вн-16',1320),('Вн-16',1020),('Вн-17',4020),('Вн-9',4060),('Вн-17',3750),('Вн-6',4130)]:
    rows.append(('без литера',m,w,4800,1,0,'остаток КП без литера, по размеру к строкам договора не подходит',None))
rows.sort(key=lambda z:(LITS.index(z[0]) if z[0] in LITS else 9,key(z[1])))
r=4
for n,(l,m,w,h,q,got,why,a1) in enumerate(rows,1):
    ws4.append([n,l,m,w,h,a1 if a1 is not None else f'=ROUND(D{r}*E{r}/1000000,2)',q,got,f'=G{r}-H{r}',f'=ROUND(I{r}*F{r},2)',why]); r+=1
ws4.cell(r,3,'ИТОГО').font=hf
for c in (7,8,9,10):
    col=L(c); ws4.cell(r,c,f'=SUM({col}4:{col}{r-1})').font=hf
for c in range(1,12): ws4.cell(r,c).fill=tf
box(ws4,3,r,1,11)
for rr in range(4,r+1):
    for c in range(1,12): ws4.cell(rr,c).alignment=W if c==11 else C
    ws4.cell(rr,6).number_format='0.00'; ws4.cell(rr,10).number_format='#,##0.00'
for i,wd in enumerate([4,11,8,9,9,10,8,13,9,12,40],1): ws4.column_dimensions[L(i)].width=wd
ws4.row_dimensions[3].height=60
ws4.page_setup.orientation='landscape'; ws4.sheet_properties.pageSetUpPr.fitToPage=True; ws4.page_setup.fitToHeight=0
# ---- Sheet 5: сверка с КП
ws5=wb.create_sheet('Сверка с КП')
ws5.append(['Сверка с КП. В спецификации площадь по КП = ширина × высота из КП (как в таблице заказчика)']); ws5['A1'].font=hf
ws5.append(['«Площадь в КП» — цифра из самого КП: она меньше, т.к. двери посчитаны отдельными КП от 13.04'])
ws5.append(['1. Позиции КП на витражи']); ws5.cell(3,1).font=hf
h5=['№','КП','№ в КП','Марка','Ширина, мм','Высота, мм','Кол-во, шт','Площадь по размерам КП, м²','Площадь в КП, м²','Разница, м²','Куда отнесено']
ws5.append(h5); hr=4
for c in range(1,12): ws5.cell(hr,c).font=hf; ws5.cell(hr,c).alignment=C; ws5.cell(hr,c).fill=fill
r=5
for n,(doc,p,m,w,h,q,a,l) in enumerate(KPP,1):
    ws5.append([n,doc,p,m,w,h,q,f'=ROUND(E{r}*F{r}*G{r}/1000000,2)',round(a,2),f'=H{r}-I{r}',{'Г←А':'А (установлен в лит. Г)','Г←Б':'Б (установлен в лит. Г)'}.get(l,l)]); r+=1
ws5.cell(r,2,'ИТОГО').font=hf
for c in (7,8,9,10): ws5.cell(r,c,f'=SUM({L(c)}5:{L(c)}{r-1})').font=hf
for c in range(1,12): ws5.cell(r,c).fill=tf
box(ws5,hr,r,1,11); e1=r
r+=2; ws5.cell(r,1,'2. Двери к витражам (КП от 13.04.2026) — в спецификацию не добавлены, для справки').font=hf; r+=1
h6=['№','КП','Литер','Марка','Ширина, мм','Высота, мм','Кол-во, шт','Площадь по размерам, м²','Площадь в КП, м²']
for c,v in enumerate(h6,1): x=ws5.cell(r,c,v); x.font=hf; x.alignment=C; x.fill=fill
hr2=r; r+=1; s2=r
for n,(doc,l,m,w,h,q,a) in enumerate(KPD,1):
    for c,v in enumerate([n,doc,l,m,w,h,q,f'=ROUND(E{r}*F{r}*G{r}/1000000,2)',a],1): ws5.cell(r,c,v)
    r+=1
ws5.cell(r,2,'ИТОГО').font=hf
for c in (7,8,9): ws5.cell(r,c,f'=SUM({L(c)}{s2}:{L(c)}{r-1})').font=hf
for c in range(1,10): ws5.cell(r,c).fill=tf
box(ws5,hr2,r,1,9); e2=r
for rr in range(hr,r+1):
    for c in range(1,12):
        x=ws5.cell(rr,c)
        if x.value is not None and rr not in (hr,hr2) and not (isinstance(x.value,str) and x.value[:1].isdigit() and '.' in x.value[:3]): x.alignment=C
    for c in (8,9,10): ws5.cell(rr,c).number_format='0.00'
for i,wd in enumerate([6,26,9,9,10,10,9,12,11,10,17],1): ws5.column_dimensions[L(i)].width=wd
for x in (hr,hr2): ws5.row_dimensions[x].height=60
ws5.page_setup.orientation='landscape'; ws5.sheet_properties.pageSetUpPr.fitToPage=True; ws5.page_setup.fitToHeight=0
# ---- Sheet 6: закрытие по ВДЦ (Приложение №2, лист «Корректировка»)
VDC=[('А', 'Вн-6', 4200, 4820, 4, '4.2*4.82*4', 18), ('А', 'Вн-8', 2600, 4820, 1, '2.6*4.82*1', 20), ('А', 'Вн-14', 2485, 4820, 1, '2.485*4.82*1', 22), ('А', 'Вн-15', 1400, 4820, 1, '1.4*4.82*1', 24), ('А', 'Вн-22', 2150, 2720, 5, '2.15*2.72*5', 26), ('А', 'Вн-23', 3540, 3020, 1, '3.54*3.02*1', 28), ('А', 'Вн-24', 2610, 3020, 1, '2.61*3.02*1', 30), ('Б', 'Вн-6', 4200, 4820, 4, '4.2*4.82*4', 42), ('Б', 'Вн-15', 1400, 4820, 1, '1.4*4.82', 44), ('Б', 'Вн-17', 4200, 4820, 1, '4.2*4.82', 46), ('Б', 'Вн-18', 2600, 4820, 1, '2.6*4.82', 48), ('Б', 'Вн-22', 2150, 2720, 5, '2.15*2.72*5', 50), ('Б', 'Вн-23', 3540, 3020, 1, '3.54*3.02*1', 52), ('Б', 'Вн-24', 2610, 3020, 1, '2.61*3.02*1', 54), ('В', 'Вн-6', 4200, 4820, 1, '4.2*4.82*1', 66), ('В', 'Вн-16', 1200, 4820, 1, '1.2*4.82*1', 68), ('В', 'Вн-13', 3200, 4820, 1, '3.2*4.82*1', 70), ('В', 'Вн-19', 3000, 4820, 2, '3*4.82*2', 72), ('В', 'Вн-22', 2150, 2720, 4, '2.15*2.72*4', 74), ('Г', 'Вн-6', 4200, 4820, 1, '4.2*4.82*1', 88), ('Г', 'Вн-10', 4160, 4820, 1, '4.16*4.82*1', 90), ('Г', 'Вн-11', 2450, 4820, 1, '2.45*4.82', 92), ('Г', 'Вн-12', 3370, 4820, 1, '3.37*4.82', 94), ('Г', 'Вн-17', 4200, 4820, 1, '4.2*4.82', 96), ('Г', 'Вн-20', 3900, 4820, 1, '3.9*4.82', 98), ('Г', 'Вн-21', 2920, 4820, 1, '2.92*4.82', 100), ('Г', 'Вн-22', 2150, 2720, 4, '2.15*2.72*4', 102), ('Д', 'Вн-6', 4200, 4820, 2, '4.2*4.82*2', 116), ('Д', 'Вн-7', 3020, 4820, 1, '3.02*4.82*1', 118), ('Д', 'Вн-9', 3870, 4820, 1, '3.87*4.82', 120), ('Д', 'Вн-17', 4200, 4820, 1, '4.2*4.82', 122), ('Ж', 'Вн-1', 4750, 3870, 1, '4.75*3.87', 136), ('Ж', 'Вн-2', 4750, 3870, 1, '4.75*3.87', 138), ('Ж', 'Вн-3', 6100, 3870, 1, '6.1*3.87', 140), ('Ж', 'Вн-4', 5400, 3470, 1, '5.4*3.47*1', 142), ('Ж', 'Вн-5', 7800, 3470, 1, '7.8*3.4', 144)]
VDCD=[('А', 'Д17*', 2410, 1300, 'П', 6, '2.41*1.3*6', 33), ('А', 'Д18', 2100, 1000, 'П', 1, '2.1*1', 35), ('А', 'Д21', 2100, 1200, 'П', 1, '2.1*1.2', 37), ('А', 'Д22', 2100, 1300, 'П', 1, '2.1*1.3', 39), ('Б', 'Д17*', 2410, 1300, 'П', 6, '2.41*1.3*6', 57), ('Б', 'Д18', 2100, 1000, 'П', 1, '2.1*1', 59), ('Б', 'Д21', 2100, 1200, 'П', 1, '2.1*1.2', 61), ('Б', 'Д22', 2100, 1300, 'П', 1, '2.1*1.3', 63), ('В', 'Д17*', 2410, 1300, 'П', 3, '2.41*1.3*3', 77), ('В', 'Д18', 2100, 1000, 'П', 1, '2.1*1', 79), ('В', 'Д35', 2410, 1300, 'Л', 1, '2.41*1.3', 81), ('В', 'Д36', 2100, 1400, 'Л', 1, '2.1*1.4', 83), ('В', 'Д37', 2100, 1400, 'Л', 1, '2.1*1.4', 85), ('Г', 'Д17*', 2410, 1300, 'П', 3, '2.41*1.3*3', 105), ('Г', 'Д18', 2100, 1000, 'П', 1, '2.1*1', 107), ('Г', 'Д22', 2100, 1300, 'П', 1, '2.1*1.3', 109), ('Г', 'Д35', 2410, 1300, 'Л', 3, '2.41*1.3*3', 111), ('Г', 'Д37', 2100, 1400, 'Л', 1, '2.1*1.4', 113), ('Д', 'Д17*', 2410, 1300, 'П', 6, '2.41*1.3*6', 125), ('Д', 'Д18', 2100, 1000, 'П', 1, '2.1*1', 127), ('Д', 'Д19', 2100, 1000, 'П', 1, '2.1*1', 129), ('Д', 'Д22', 2100, 1300, 'П', 1, '2.1*1.3', 131), ('Д', 'Д35', 2410, 1300, 'Л', 1, '2.41*1.3', 133)]
# двери по КП (13.04): тип — по марке двери у витража на чертеже (Вн-16 — Д35, остальные — Д17*)
# КП №15 на 11 дверей Вн-22 распределён так же, как сами Вн-22 по позициям ВДЦ: А5, Б5, В1
V22={'А':5,'Б':5,'В':1}
KD=collections.Counter(); KDL=[]
for doc,l,host,w,h,q,a in KPD:
    dm='Д35' if host=='Вн-16' else 'Д17*'
    if host=='Вн-22':
        for ll,qq in V22.items():
            KD[(ll,dm)]+=qq; KDL.append((doc,ll,host,dm,w,h,qq,round(a*qq/q,2)))
    else:
        KD[(l,dm)]+=q; KDL.append((doc,l,host,dm,w,h,q,a))
# распределение дверей КП по строкам ВДЦ: сначала свой литер, излишек — в свободные строки других литеров
dclose={}; dsrc=collections.defaultdict(list); surplus=collections.defaultdict(list)
for l,dm,h_,w_,side,q,f,rr in VDCD:
    own=min(q,KD[(l,dm)]); dclose[(l,dm)]=own
    if own: dsrc[(l,dm)].append(f'{own} — КП лит. {l}')
    if KD[(l,dm)]>q: surplus[dm].append([l,KD[(l,dm)]-q])
for l in LITS:
    for dm in {d for _,d in KD}:
        if KD[(l,dm)] and not any(x[0]==l and x[1]==dm for x in VDCD): surplus[dm].append([l,KD[(l,dm)]])
for l,dm,h_,w_,side,q,f,rr in VDCD:
    need=q-dclose[(l,dm)]
    for sp in surplus[dm]:
        if need<=0: break
        t=min(need,sp[1])
        if t: sp[1]-=t; need-=t; dclose[(l,dm)]+=t; dsrc[(l,dm)].append(f'{t} — излишек КП лит. {sp[0]}')
dleft=[(dm,l,n) for dm,v in surplus.items() for l,n in v if n]

ws6=wb.create_sheet('Закрытие по ВДЦ',1)
ws6.append(['Закрытие по ВДЦ (Приложение № 2, лист «Корректировка»): к закрытию — меньшее из ВДЦ и КП, площадь по размерам ВДЦ']); ws6['A1'].font=Font(bold=True,size=12)
ws6.append(['Строка ВДЦ — номер строки «Изготовление и монтаж» на листе «Корректировка». Жёлтым — количество не совпадает с ВДЦ'])
yel=PatternFill('solid',fgColor='FFF2CC')
h=['№','Литер','Марка','Размер по ВДЦ, мм','Строка ВДЦ','По ВДЦ, шт','По ВДЦ, м²','По КП / спецификации, шт','К закрытию, шт','К закрытию, м²','Не закрыто по ВДЦ, шт','Сверх ВДЦ, шт','Примечание']
NC=len(h)
def head(r):
    for c,v in enumerate(h,1): x=ws6.cell(r,c,v); x.font=hf; x.alignment=C; x.fill=fill
    ws6.row_dimensions[r].height=62
def fmt(r1,r2,notecol=13):
    box(ws6,r1-1,r2,1,NC)
    for rr in range(r1,r2+1):
        for c in range(1,NC+1): ws6.cell(rr,c).alignment=W if c==notecol else C
        for c in (7,10): ws6.cell(rr,c).number_format='#,##0.00'
r=4; ws6.cell(r,1,'1. Витражи').font=hf; r+=1; head(r); r+=1; s1=r
vmap={(l,m):(w,hh,q,f,rr) for l,m,w,hh,q,f,rr in VDC}
keys=list(vmap)+[(l,m) for l in LITS for m in marks if cnt[(m,l)] and (l,m) not in vmap]
order=['А','Б','В','Г','Д','Ж']
keys.sort(key=lambda k:(order.index(k[0]),key(k[1])))
VNOTE={('А','Вн-6'):'1 шт. установлена в лит. Г (КП 1-18 №6), объём записан в лит. А',
       ('Б','Вн-6'):'1 шт. установлена в лит. Г (КП 1-18 №8), объём записан в лит. Б',
       ('Г','Вн-6'):'на чертеже 3 шт., 2 из них закрываются по лит. А и Б',
       ('В','Вн-16'):'по чертежу и КП 3 шт., в ВДЦ 1',('В','Вн-17'):'в ВДЦ лит. В нет',
       ('Ж','Вн-5'):'в ВДЦ площадь =7.8*3.4 (26,52), а в названии 7800х3470 (27,07)'}
V22N='КП на 11 шт. распределён по позициям ВДЦ: А5, Б5, В1'
n=1
for l,m in keys:
    q_s=cnt[(m,l)] if l in LITS else 0
    if (l,m) in vmap:
        w,hh,q,f,rr=vmap[(l,m)]
        row=[n,l,m,f'{w}×{hh}',rr,q,'='+f,q_s or None,f'=MIN(F{r},N(H{r}))',f'=ROUND(G{r}/F{r}*I{r},2)',f'=F{r}-I{r}',f'=N(H{r})-I{r}']
    else:
        q=0; row=[n,l,m,None,None,None,None,q_s,0,0,0,f'=N(H{r})']
    st=VNOTE.get((l,m)) or (V22N if m=='Вн-22' else None) or ('совпадает' if q==q_s else ('нет в КП' if q_s==0 else ('в ВДЦ нет' if q==0 else f'в ВДЦ {q}, по КП {q_s}')))
    ws6.append(row+[st])
    if q!=q_s and (l,m)!=('Г','Вн-6'):
        for c in range(1,NC+1): ws6.cell(r,c).fill=yel
    n+=1; r+=1
ws6.cell(r,3,'Итого витражи').font=hf
for c in range(6,13): ws6.cell(r,c,f'=SUM({L(c)}{s1}:{L(c)}{r-1})').font=hf
for c in range(1,NC+1): ws6.cell(r,c).fill=tf
fmt(s1,r); t1=r
r+=2; ws6.cell(r,1,'2. Двери (КП от 13.04.2026; марка — по чертежу у витража). Излишек КП одного литера закрывает свободные строки ВДЦ других литеров').font=hf; r+=1
head(r); ws6.cell(r,8,'По КП (свой литер), шт'); r+=1; s2=r
dorder={'Д17*':0,'Д18':1,'Д19':2,'Д21':3,'Д22':4,'Д35':5,'Д36':6,'Д37':7}
n=1
for l,dm,h_,w_,side,q,f,rr in sorted(VDCD,key=lambda x:(order.index(x[0]),dorder[x[1]])):
    k=(l,dm); cl=dclose.get(k,0)
    note='; '.join(dsrc[k]) if dsrc[k] else 'нет в КП'
    ws6.append([n,l,f'{dm} ({side})',f'{h_}×{w_}',rr,q,'='+f,KD[k] or None,cl,f'=ROUND(G{r}/F{r}*I{r},2)',f'=F{r}-I{r}',f'=MAX(0,N(H{r})-F{r})',note])
    if cl!=q:
        for c in range(1,NC+1): ws6.cell(r,c).fill=yel
    n+=1; r+=1
for dm,l,left in dleft:
    ws6.append([n,l,dm,None,None,None,None,None,0,0,0,None,f'{left} шт. из КП лит. {l} — свободных строк ВДЦ {dm} больше нет, закрыть нельзя'])
    for c in range(1,NC+1): ws6.cell(r,c).fill=yel
    n+=1; r+=1
ws6.cell(r,3,'Итого двери').font=hf
for c in range(6,13): ws6.cell(r,c,f'=SUM({L(c)}{s2}:{L(c)}{r-1})').font=hf
for c in range(1,NC+1): ws6.cell(r,c).fill=tf
fmt(s2,r); t2=r
r+=2
ws6.cell(r,3,'ВСЕГО к закрытию (витражи + двери)').font=hf
ws6.merge_cells(start_row=r,start_column=3,end_row=r,end_column=8)
ws6.cell(r,9,f'=I{t1}+I{t2}').font=hf; ws6.cell(r,10,f'=J{t1}+J{t2}').font=hf; ws6.cell(r,10).number_format='#,##0.00'
for c in range(1,NC+1): ws6.cell(r,c).fill=tf
box(ws6,r,r,1,NC)
for i,wd in enumerate([4,6,10,12,8,8,10,12,10,11,11,9,44],1): ws6.column_dimensions[L(i)].width=wd
ws6.freeze_panes='D4'
ws6.page_setup.orientation='landscape'; ws6.sheet_properties.pageSetUpPr.fitToPage=True; ws6.page_setup.fitToHeight=0

# ---- Sheet 7: двери по КП
ws7=wb.create_sheet('Двери',3)
ws7.append(['Двери в составе витражей по КП от 13.04.2026']); ws7['A1'].font=Font(bold=True,size=12)
ws7.append(['Марка двери — по чертежу у витража; двери Вн-22 (КП №15, 11 шт.) распределены как сами Вн-22: А5, Б5, В1'])
h7=['№','Литер','Витраж','Марка двери','КП','Ширина, мм','Высота, мм','Кол-во, шт','Площадь по размерам КП, м²','Площадь в КП, м²']
ws7.append(h7)
for c in range(1,11): x=ws7.cell(3,c); x.font=hf; x.alignment=C; x.fill=fill
r=4
for n,(doc,l,host,dm,w,hh,q,a) in enumerate(sorted(KDL,key=lambda x:(order.index(x[1]),key(x[2]))),1):
    ws7.append([n,l,host,dm,doc,w,hh,q,f'=ROUND(F{r}*G{r}*H{r}/1000000,2)',a]); r+=1
ws7.cell(r,3,'ИТОГО').font=hf
for c in (8,9,10): ws7.cell(r,c,f'=SUM({L(c)}4:{L(c)}{r-1})').font=hf
for c in range(1,11): ws7.cell(r,c).fill=tf
box(ws7,3,r,1,10)
for rr in range(4,r+1):
    for c in range(1,11): ws7.cell(rr,c).alignment=C
    for c in (9,10): ws7.cell(rr,c).number_format='0.00'
for i,wd in enumerate([4,6,8,9,11,9,9,8,12,10],1): ws7.column_dimensions[L(i)].width=wd
ws7.row_dimensions[3].height=60
wb.calculation.fullCalcOnLoad=True
wb.save('Спецификация_витражей.xlsx')
